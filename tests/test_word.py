"""Tester for verktøyene på Word-dokumenter (.docx). Kjører uten Claude – bare ren Python."""

from pathlib import Path

from docx import Document

from docxreader import tools
from docxreader.tools import document_outline, read_document, read_section, search_document

SAMPLES = Path(__file__).parent.parent / "samples"
SAMPLE_DOCX = SAMPLES / "prosjektplan.docx"
LARGE_DOCX = SAMPLES / "arsrapport.docx"


def test_docx_headings_become_markdown():
    result = read_document.invoke({"path": str(SAMPLE_DOCX)})
    assert "# Prosjektplan: Ny nettbutikk" in result
    assert "\n## Budsjett\n" in result
    assert "\n### Fase 1: Design\n" in result


def test_docx_lists():
    result = read_document.invoke({"path": str(SAMPLE_DOCX)})
    assert "- Lansere ny nettbutikk i november\n- Halvere tiden" in result
    assert "1. Sett opp nettbutikkplattform\n2. Integrer betaling\n3. Test med ekte kunder" in result


def test_docx_table_is_markdown_and_in_order():
    result = read_document.invoke({"path": str(SAMPLE_DOCX)})
    assert "| Post | Beløp (kr) | Ansvarlig |" in result
    assert "| Markedsføring | 150 000 | Per Æsøy |" in result
    # Tabellen skal stå MELLOM avsnittet før og avsnittet etter den.
    before = result.index("Tabellen under viser")
    table = result.index("| Post |")
    after = result.index("Totalt budsjett")
    assert before < table < after


def test_docx_header_and_footer():
    result = read_document.invoke({"path": str(SAMPLE_DOCX)})
    assert "Topptekst: Fjellbekk AS – Internt" in result
    assert "Bunntekst: Konfidensielt" in result


def test_docx_numbering_restarts_after_other_content(tmp_path):
    doc = Document()
    doc.add_paragraph("A", style="List Number")
    doc.add_paragraph("B", style="List Number")
    doc.add_paragraph("Mellomtekst")
    doc.add_paragraph("C", style="List Number")
    f = tmp_path / "liste.docx"
    doc.save(f)

    result = read_document.invoke({"path": str(f)})
    assert "1. A\n2. B\n\nMellomtekst\n\n1. C" in result


def test_docx_table_cell_with_pipe_and_newline(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.rows[0].cells[0].text = "Kolonne"
    table.rows[1].cells[0].text = "a | b\nc"
    f = tmp_path / "tabell.docx"
    doc.save(f)

    result = read_document.invoke({"path": str(f)})
    assert "| a \\| b c |" in result


def test_docx_empty_document(tmp_path):
    f = tmp_path / "tom.docx"
    Document().save(f)
    assert "ingen tekst" in read_document.invoke({"path": str(f)})


def test_docx_missing_file_returns_error_text():
    assert read_document.invoke({"path": "finnes_ikke.docx"}).startswith("Feil:")


def test_docx_corrupt_file_returns_error_text(tmp_path):
    bad = tmp_path / "ødelagt.docx"
    bad.write_bytes(b"dette er ikke en ekte docx")
    result = read_document.invoke({"path": str(bad)})
    assert result.startswith("Feil:")
    assert "ikke et gyldig Word-dokument" in result


# --- Store dokumenter: innholdsfortegnelse, seksjoner og grense -------------


def _chapters_doc(tmp_path) -> str:
    """Lite dokument med samme underoverskrift ("Status") i to kapitler."""
    doc = Document()
    doc.add_heading("Rapport", level=0)
    for chapter in ["Nord", "Sør"]:
        doc.add_heading(chapter, level=1)
        doc.add_paragraph(f"Innledning for {chapter}.")
        doc.add_heading("Status", level=2)
        doc.add_paragraph(f"Status i {chapter} er god.")
    f = tmp_path / "kapitler.docx"
    doc.save(f)
    return str(f)


def test_outline_lists_headings_with_sizes():
    result = document_outline.invoke({"path": str(SAMPLE_DOCX)})
    assert "## Budsjett  (~33 ord, 1 tabell)" in result
    assert "### Fase 2: Utvikling" in result
    # Bare overskrifter – ikke selve teksten.
    assert "Kari Nordmann" not in result


def test_outline_without_headings(tmp_path):
    doc = Document()
    doc.add_paragraph("Bare tekst.")
    f = tmp_path / "flat.docx"
    doc.save(f)
    assert "ingen overskrifter" in document_outline.invoke({"path": str(f)})


def test_section_includes_content_until_next_heading():
    result = read_section.invoke({"path": str(SAMPLE_DOCX), "heading": "Budsjett"})
    assert result.startswith("## Budsjett")
    assert "| Utvikling | 850 000 | Kari Nordmann |" in result
    assert "Totalt budsjett" in result
    assert "Fremdrift" not in result


def test_section_includes_subsections():
    result = read_section.invoke({"path": str(SAMPLE_DOCX), "heading": "Fremdrift"})
    assert "### Fase 1: Design" in result
    assert "3. Test med ekte kunder" in result


def test_section_case_insensitive_and_partial_match():
    result = read_section.invoke({"path": str(SAMPLE_DOCX), "heading": "fase 2"})
    assert result.startswith("### Fase 2: Utvikling")


def test_section_not_found_lists_available_headings():
    result = read_section.invoke({"path": str(SAMPLE_DOCX), "heading": "Lønn"})
    assert result.startswith("Feil:")
    assert "Fremdrift > Fase 1: Design" in result


def test_section_ambiguous_heading_asks_for_path(tmp_path):
    result = read_section.invoke({"path": _chapters_doc(tmp_path), "heading": "Status"})
    assert result.startswith("Feil:")
    assert "Nord > Status" in result and "Sør > Status" in result


def test_section_path_picks_the_right_one(tmp_path):
    result = read_section.invoke({"path": _chapters_doc(tmp_path), "heading": "Sør > Status"})
    assert "Status i Sør er god." in result
    assert "Nord" not in result


def test_section_missing_file_returns_error_text():
    result = read_section.invoke({"path": "finnes_ikke.docx", "heading": "x"})
    assert result.startswith("Feil:")


def test_read_document_returns_outline_when_too_large(tmp_path, monkeypatch):
    # monkeypatch endrer MAX_WORDS bare under denne testen.
    monkeypatch.setattr(tools, "MAX_WORDS", 10)
    result = read_document.invoke({"path": _chapters_doc(tmp_path)})
    assert "Dokumentet er stort" in result
    assert "read_section" in result
    assert "### Status" in result
    assert "er god" not in result  # selve teksten skal ikke være med


def test_large_section_returns_sub_outline(tmp_path, monkeypatch):
    monkeypatch.setattr(tools, "MAX_WORDS", 5)
    result = read_section.invoke({"path": _chapters_doc(tmp_path), "heading": "Nord"})
    assert "er stor" in result
    assert "### Status" in result
    assert "er god" not in result


def test_large_sample_read_document_gives_outline():
    result = read_document.invoke({"path": str(LARGE_DOCX)})
    assert result.startswith("Dokumentet er stort")
    assert "## Region Nord" in result
    assert len(result) < 3000  # mye mindre enn hele dokumentet


def test_large_sample_planted_facts_are_reachable():
    status = read_section.invoke({"path": str(LARGE_DOCX), "heading": "Region Nord > Status"})
    assert "Ingrid Solberg" in status
    it = read_section.invoke({"path": str(LARGE_DOCX), "heading": "IT og sikkerhet > Utfordringer"})
    assert "3 alvorlige sikkerhetshendelser" in it
    tall = read_section.invoke({"path": str(LARGE_DOCX), "heading": "Økonomi > Nøkkeltall"})
    assert "| Driftsresultat | 31,2 MNOK | 38,5 MNOK |" in tall


# --- search_document -------------------------------------------------------------


def _search(path, query) -> str:
    return search_document.invoke({"path": str(path), "query": query})


def test_search_returns_sentence_with_heading_path():
    result = _search(LARGE_DOCX, "kontaktperson")
    assert result.startswith("1 treff")
    assert "[Region Nord > Status]" in result
    assert "- Kontaktperson for Region Nord er Ingrid Solberg." in result
    # Bare setningen, ikke hele avsnittet rundt.
    assert len(result) < 200


def test_search_matches_part_of_word_and_ignores_case():
    result = _search(LARGE_DOCX, "SIKKERHETSHENDELSE")
    assert "3 alvorlige sikkerhetshendelser" in result
    assert "[IT og sikkerhet > Utfordringer]" in result


def test_search_table_row_comes_with_header_row():
    result = _search(LARGE_DOCX, "driftsresultat")
    assert "[Økonomi > Nøkkeltall]" in result
    assert "| Nøkkeltall | 2025 | 2026 |" in result
    assert "| Driftsresultat | 31,2 MNOK | 38,5 MNOK |" in result
    # Andre rader i tabellen skal ikke være med.
    assert "| Omsetning |" not in result


def test_search_all_words_must_be_in_same_sentence():
    assert "Ingrid Solberg" in _search(LARGE_DOCX, "kontaktperson nord")
    assert _search(LARGE_DOCX, "kontaktperson sør").startswith("Ingen treff")


def test_search_list_items_and_table_cells_in_small_doc():
    assert "3. Test med ekte kunder" in _search(SAMPLE_DOCX, "ekte kunder")
    result = _search(SAMPLE_DOCX, "Nordmann")
    assert "[Budsjett]" in result  # tittelen er ikke med i stien
    assert "| Utvikling | 850 000 | Kari Nordmann |" in result


def test_search_limits_number_of_hits(monkeypatch):
    monkeypatch.setattr(tools, "MAX_SEARCH_HITS", 3)
    result = _search(LARGE_DOCX, "kundene")  # vanlig ord i fylltekst
    assert "Viser 3 av" in result
    assert result.count("\n- ") == 3


def test_search_no_hits_gives_tips():
    result = _search(LARGE_DOCX, "lønnsoppgjør")
    assert result.startswith("Ingen treff")
    assert "document_outline" in result


def test_search_empty_query_returns_error_text():
    assert _search(SAMPLE_DOCX, "   ").startswith("Feil:")


def test_search_missing_file_returns_error_text():
    assert _search("finnes_ikke.docx", "x").startswith("Feil:")
