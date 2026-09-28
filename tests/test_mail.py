"""Tester for e-post (.eml og .msg). Testfilen lages av scripts/make_sample_eml.py."""

import datetime
import io
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from extract_msg.enums import AttachmentType

from docxreader.readers import mail
from docxreader.tools import document_outline, read_document, read_section, search_document

SAMPLE = str(Path(__file__).parent.parent / "samples" / "tilbud.eml")


def _section(heading: str, path: str = SAMPLE) -> str:
    return read_section.invoke({"path": path, "heading": heading})


# --- .eml: den ekte testfilen -------------------------------------------------


def test_eml_header_is_decoded_into_meta():
    result = read_document.invoke({"path": SAMPLE})
    assert result.startswith(
        "E-post\n"
        "Fra: Kari Nordmann <kari@fjellbekk.example>\n"
        "Til: Ola Hansen <ola@fjellbekk.example>\n"
        "Kopi: Per Æsøy <per@fjellbekk.example>\n"
        "Dato: 2026-09-28 09:15\n"
        "Emne: Re: Tilbud på servere – Q4\n"
        "Vedlegg: tilbud.xlsx, avtaleutkast.docx, skisse.png, Tilbud 2026-117.eml"
    )


def test_eml_prefers_plain_text_and_skips_inline_logo():
    message = _section("Melding")
    assert "Kan du se over betalingsbetingelsene før fredag?" in message  # står bare i ren tekst
    assert "logo" not in read_document.invoke({"path": SAMPLE})


def test_eml_thread_is_split_into_quoted_messages():
    outline = document_outline.invoke({"path": SAMPLE})
    assert "## Melding  (~" in outline
    assert "## Tidligere melding 1  (~" in outline
    assert "## Tidligere melding 2  (~" in outline
    first = _section("Tidligere melding 1")
    assert "Den 25. sep. 2026 kl. 14:02 skrev Ola Hansen <ola@fjellbekk.example>:\n\nHei Kari," in first
    assert not any(line.startswith(">") for line in first.splitlines())  # sitat-tegnet er fjernet
    assert "Opprinnelig melding" not in first


def test_eml_outlook_header_in_thread_becomes_list():
    second = _section("Tidligere melding 2")
    assert second.startswith("## Tidligere melding 2\n\n-----Opprinnelig melding-----\n\n- Fra: Per Æsøy")
    assert "- Emne: Leverandørmøte\n\nMøtet med leverandøren" in second


def test_eml_planted_fact_deep_in_thread_is_searchable():
    result = search_document.invoke({"path": SAMPLE, "query": "fjorden"})
    assert "[Tidligere melding 2]" in result
    assert "torsdag 8. oktober kl. 10" in result


def test_eml_attachments_are_read_by_their_own_reader():
    xlsx = _section("Vedlegg: tilbud.xlsx")
    assert "- Tilbud: 2 rader; kolonner: Leverandør, Vare, Antall, Pris per stk" in xlsx
    assert "### Ark: Tilbud" in xlsx  # vedleggets øverste nivå er ###
    assert "| 2 | Nordic Data AS | Server R750 | 4 | 45000 |" in xlsx
    docx = _section("Vedlegg: avtaleutkast.docx")
    assert "### Avtaleutkast" in docx and "#### Betalingsbetingelser" in docx


def test_eml_section_inside_attachment_by_path():
    result = _section("avtaleutkast > Betalingsbetingelser")
    assert result.startswith("#### Betalingsbetingelser")
    assert "30 dager etter levering" in result


def test_eml_image_attachment_points_to_view_file():
    assert 'PDF/bilde – se innholdet med view_file og attachment="skisse.png")' in _section("skisse.png")


def test_eml_forwarded_mail_attachment():
    result = _section("Vedlegg: Tilbud 2026-117.eml")
    assert "Fra: Salg <salg@nordicdata.example>" in result
    assert "### Melding" in result
    assert "Tilbudet gjelder til 15. oktober 2026." in result


# --- .eml: laget i testen ---------------------------------------------------


def _write_eml(tmp_path, message: EmailMessage, name="a.eml") -> str:
    f = tmp_path / name
    f.write_bytes(message.as_bytes())
    return str(f)


def _message(body: str = "Hei.", subject: str = "Test") -> EmailMessage:
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = "a@x.no", "b@x.no", subject
    m.set_content(body)
    return m


def test_html_only_mail_is_converted_to_text(tmp_path):
    m = EmailMessage()
    m["From"], m["Subject"] = "a@x.no", "HTML"
    m.set_content(
        "<html><head><style>p{}</style></head><body><h1>Status</h1><p>Alt går <b>bra</b>.</p>"
        "<ul><li>Én</li><li>To</li></ul><table><tr><th>A</th><th>B</th></tr>"
        "<tr><td>1</td><td>2</td></tr></table></body></html>",
        subtype="html",
    )
    result = read_document.invoke({"path": _write_eml(tmp_path, m)})
    assert "## Melding\n\n### Status\n\nAlt går bra." in result  # <h1> flyttet under Melding
    assert "- Én\n- To" in result
    assert "| A | B |\n| --- | --- |\n| 1 | 2 |" in result
    assert "p{}" not in result


@pytest.mark.parametrize("marker", [
    "On Mon, 21 Sep 2026 at 08:30, Ola Hansen <ola@x.no> wrote:",
    "---------- Forwarded message ---------",
    "-----Original Message-----",
])
def test_other_reply_markers(tmp_path, marker):
    f = _write_eml(tmp_path, _message(f"Svar.\n\n{marker}\n> Gammel tekst."))
    assert "Gammel tekst." in _section("Tidligere melding 1", f)
    assert "Gammel" not in _section("Melding", f)


def test_outlook_header_without_separator_line(tmp_path):
    body = "Ok!\n\nFra: Per <per@x.no>\nSendt: 20. september 2026\nEmne: Hei\n\nGammel tekst."
    result = _section("Tidligere melding 1", _write_eml(tmp_path, _message(body)))
    assert "- Fra: Per <per@x.no>\n- Sendt: 20. september 2026" in result


def test_broken_attachment_gives_note_not_error(tmp_path):
    m = _message()
    m.add_attachment(b"ikke et word-dokument", maintype="application", subtype="octet-stream", filename="rapport.docx")
    result = read_document.invoke({"path": _write_eml(tmp_path, m)})
    assert "## Vedlegg: rapport.docx" in result
    assert "kunne ikke leses: 'rapport.docx' er ikke et gyldig Word-dokument" in result


def test_attachment_filename_cannot_escape_temp_folder(tmp_path):
    m = _message()
    buffer = io.BytesIO()
    Document().save(buffer)
    m.add_attachment(buffer.getvalue(), maintype="application", subtype="octet-stream", filename="../../hemmelig.docx")
    assert "## Vedlegg: hemmelig.docx" in read_document.invoke({"path": _write_eml(tmp_path, m)})


def test_deeply_nested_forwarded_mails_stop(tmp_path):
    inner = _message("Innerst.", "Nivå 0")
    for level in range(1, 6):
        outer = _message(f"Nivå {level}.", f"Nivå {level}")
        outer.add_attachment(inner)
        inner = outer
    result = read_document.invoke({"path": _write_eml(tmp_path, inner)})
    assert "for dypt nøstet" in result
    assert "Innerst." not in result


def test_text_file_renamed_to_eml_is_rejected(tmp_path):
    f = tmp_path / "notat.eml"
    f.write_text("Bare litt tekst uten e-posthode.")
    result = read_document.invoke({"path": str(f)})
    assert result.startswith("Feil:") and "ser ikke ut som en e-post" in result


def test_binary_file_as_eml_is_rejected(tmp_path):
    f = tmp_path / "bilde.eml"
    f.write_bytes(b"\x89PNG\x00\x00\x00")
    assert read_document.invoke({"path": str(f)}).startswith("Feil:")


# --- .msg (Outlook) ----------------------------------------------------------
# extract-msg kan ikke LAGE .msg-filer, så vi bytter ut openMsg med en funksjon
# som gir et falskt Outlook-objekt med de samme feltene. Da testes vår kode
# (oversettelsen til Mail), ikke biblioteket.


def _fake_msg(subject="Møte", body="Hei!\n\nVelkommen.", attachments=(), html=None):
    return SimpleNamespace(
        sender="Kari <kari@x.no>", to="Ola <ola@x.no>", cc=None,
        date=datetime.datetime(2026, 9, 28, 9, 15), subject=subject,
        body=body, htmlBody=html, attachments=list(attachments), close=lambda: None,
    )


class _FakeAttachment(SimpleNamespace):
    def __or__(self, changes: dict):  # a | {"felt": verdi} -> kopi med endrede felt
        return _FakeAttachment(**{**vars(self), **changes})


def _fake_attachment(name, data, kind=AttachmentType.DATA):
    return _FakeAttachment(longFilename=name, shortFilename=None, name=name, data=data,
                           type=kind, mimetype="", hidden=False, contentId=None)


def _load_fake(tmp_path, monkeypatch, fake) -> str:
    monkeypatch.setattr(mail.extract_msg, "openMsg", lambda path: fake)
    f = tmp_path / "melding.msg"
    f.write_bytes(b"")  # innholdet leses ikke; openMsg er byttet ut
    return read_document.invoke({"path": str(f)})


def test_msg_fields_body_and_attachments(tmp_path, monkeypatch):
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Innhold i vedlegget.")
    doc.save(buffer)
    forwarded = _fake_msg(subject="Gammel", body="Videresendt tekst.")
    fake = _fake_msg(attachments=[
        _fake_attachment("notat.docx", buffer.getvalue()),
        _fake_attachment("Gammel", forwarded, kind=AttachmentType.MSG),
        _fake_attachment("logo.png", b"x") | {"mimetype": "image/png", "contentId": "logo1"},  # innebygd
        _fake_attachment("lenke", None, kind=AttachmentType.WEB),
    ])
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert result.startswith("E-post\nFra: Kari <kari@x.no>\nTil: Ola <ola@x.no>\nDato: 2026-09-28 09:15\nEmne: Møte")
    assert "Vedlegg: notat.docx, Gammel.msg, lenke" in result  # innebygd logo er ikke med
    assert "## Melding\n\nHei!\n\nVelkommen." in result
    assert "## Vedlegg: notat.docx\n\nInnhold i vedlegget." in result
    assert "## Vedlegg: Gammel.msg" in result and "Videresendt tekst." in result
    assert "## Vedlegg: lenke\n\n(vedleggstypen kan ikke leses)" in result


def test_msg_html_body_when_no_plain_text(tmp_path, monkeypatch):
    fake = _fake_msg(body="", html="<p>Fra <b>HTML</b>.</p>".encode())
    assert "## Melding\n\nFra HTML." in _load_fake(tmp_path, monkeypatch, fake)


def test_msg_corrupt_file_returns_error_text(tmp_path):
    bad = tmp_path / "ødelagt.msg"
    bad.write_bytes(b"ikke en outlook-fil")
    result = read_document.invoke({"path": str(bad)})
    assert result.startswith("Feil:")
    assert "ikke en gyldig Outlook-fil" in result


# Regresjonstester: feil funnet ved å lese ekte Outlook-filer (eksempelfilene
# til extract-msg; ikke lagt i repoet pga. lisens). De falske objektene
# etterligner det de ekte filene inneholdt.


class _BrokenMsg(SimpleNamespace):
    @property
    def to(self):  # som unicode-header.msg: feltet har ødelagt tegnkoding
        raise UnicodeDecodeError("gb2312", b"\xea", 0, 1, "illegal multibyte sequence")


def _recipient(name, address, kind):
    return SimpleNamespace(name=name, email=address, type=kind)


def test_msg_null_characters_are_removed(tmp_path, monkeypatch):
    fake = _fake_msg(subject="Test\x00\x00\x00", body="Tekst.\r\n\x00\x00\x00")
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert "\x00" not in result
    assert "Emne: Test\n" in result


def test_msg_recipients_come_from_recipient_list(tmp_path, monkeypatch):
    from extract_msg.enums import RecipientType

    fake = _fake_msg()
    fake.recipients = [
        _recipient("Alice\x00", "alice@x.no\x00", RecipientType.TO),
        _recipient("Dave", "dave@x.no", RecipientType.CC),
        _recipient("Carol", "carol@x.no", RecipientType.TO),
        _recipient("Alice", "alice@x.no", RecipientType.TO),  # duplikat
    ]
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert "Til: Alice <alice@x.no>, Carol <carol@x.no>\nKopi: Dave <dave@x.no>" in result


def test_msg_field_that_cannot_be_decoded_does_not_crash(tmp_path, monkeypatch):
    fake = _BrokenMsg(**{k: v for k, v in vars(_fake_msg()).items() if k != "to"})
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert result.startswith("E-post\nFra: Kari <kari@x.no>\nDato:")  # Til mangler, resten er med
    assert "Velkommen." in result


# --- Andre Outlook-elementer (classType) ----------------------------------------


def test_msg_task_fields(tmp_path, monkeypatch):
    from extract_msg.enums import TaskStatus

    fake = _fake_msg(subject="Bestille servere", body="Husk rabattkoden.")
    fake.classType = "IPM.Task\x00"
    fake.taskDueDate = datetime.datetime(2026, 10, 15, 0, 0)
    fake.taskStatus = TaskStatus.IN_PROGRESS
    fake.percentComplete = 0.5
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert result.startswith(
        "Oppgave (Outlook)\nTittel: Bestille servere\nFrist: 2026-10-15 00:00\nStatus: pågår\nFullført: 50 %"
    )
    assert "Fra:" not in result  # e-posthode gir ikke mening for en oppgave
    assert "## Notater\n\nHusk rabattkoden." in result


def test_msg_meeting_request_by_class_type(tmp_path, monkeypatch):
    fake = _fake_msg(subject="Styremøte")
    fake.classType = "IPM.Schedule.Meeting.Request"
    fake.startDate = datetime.datetime(2026, 11, 14, 9, 0)
    fake.location = "Ålesund"
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert result.startswith("Møteinnkalling (Outlook)\nFra: Kari <kari@x.no>")
    assert "Emne: Styremøte\nStart: 2026-11-14 09:00\nSted: Ålesund" in result


def test_msg_sticky_note_without_text_is_not_empty(tmp_path, monkeypatch):
    fake = _fake_msg(body="")
    fake.classType = "IPM.StickyNote"
    result = _load_fake(tmp_path, monkeypatch, fake)
    assert result.startswith("Notat (Outlook)\nOpprettet: 2026-09-28 09:15")
    assert "## Notater" not in result
