"""Tester for Excel (.xlsx). Testfilen lages av scripts/make_sample_xlsx.py."""

from pathlib import Path

from openpyxl import Workbook

from docxreader import tables
from docxreader.tools import document_outline, read_document, read_section, search_document

SAMPLE = Path(__file__).parent.parent / "samples" / "budsjett.xlsx"


def _section(heading: str) -> str:
    return read_section.invoke({"path": str(SAMPLE), "heading": heading})


def test_sheets_become_sections_and_meta_lists_sheets():
    result = read_document.invoke({"path": str(SAMPLE)})
    assert result.startswith(
        "Arbeidsbok med 4 ark:\n"
        "- Oversikt: 3 rader; kolonner: Avdeling, Budsjett, Forbruk, Rest, Sist oppdatert\n"
        "- Transaksjoner: 250 rader; kolonner: Dato, Leverandør, Beløp (kr), Avdeling\n"
        "- Tomt: tomt\n"
        "- Hjelpetall: 1 rader; kolonner: Nøkkel, Verdi"
    )
    assert "## Ark: Oversikt" in result
    assert "## Ark: Tomt" not in result  # tomme ark er bare nevnt i metateksten


def test_table_has_row_numbers_and_clean_values():
    result = _section("Oversikt")
    assert "| Rad | Avdeling | Budsjett | Forbruk | Rest | Sist oppdatert |" in result
    # Heltall uten ".0", dato uten klokkeslett, desimaltall beholdes.
    assert "| 2 | Salg | 1200000 | 950000 | =B2-C2 | 2026-09-01 |" in result
    assert "| 3 | IT | 800000 | 1150000.5 |" in result


def test_rows_in_sheet_order_regardless_of_write_order():
    result = _section("Oversikt")
    assert result.index("| 2 | Salg") < result.index("| 3 | IT") < result.index("| 4 | HR")


def test_uncalculated_formula_is_shown_as_formula():
    # Filen er aldri åpnet i Excel, så formelen har ikke noe lagret resultat.
    assert "=B3-C3" in _section("Oversikt")


def test_cell_comment_is_included_and_searchable():
    assert "- C3: Overforbruket skyldes engangskjøp av servere fra Nordic Data AS" in _section("Oversikt")
    result = search_document.invoke({"path": str(SAMPLE), "query": "servere"})
    assert "[Ark: Oversikt]" in result


def test_large_sheet_is_split_in_row_parts():
    outline = document_outline.invoke({"path": str(SAMPLE)})
    assert "## Ark: Transaksjoner  (~" in outline
    assert "### Rad 2–101" in outline
    assert "### Rad 102–201" in outline
    assert "### Rad 202–251" in outline


def test_each_part_repeats_header_row():
    part = _section("Transaksjoner > Rad 202–251")
    assert part.startswith("### Rad 202–251")
    assert "| Rad | Dato | Leverandør | Beløp (kr) | Avdeling |" in part
    assert "| 251 |" in part and "| 201 |" not in part


def test_planted_row_found_by_search_with_header_and_row_number():
    result = search_document.invoke({"path": str(SAMPLE), "query": "fjellsikring"})
    assert result.startswith("1 treff")
    assert "[Ark: Transaksjoner > Rad 102–201]" in result
    assert "| Rad | Dato | Leverandør | Beløp (kr) | Avdeling |" in result
    assert "| 180 | 2026-06-12 | Fjellsikring AS | 1250000 | IT |" in result


def test_hidden_sheet_is_marked():
    assert _section("Hjelpetall").startswith("## Ark: Hjelpetall (skjult)")


def test_trailing_empty_columns_and_empty_rows_are_dropped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Navn", "Alder"
    ws["A3"], ws["B3"] = "Kari", 41  # rad 2 er tom
    ws["F10"] = ""  # "brukt" celle langt ute, uten innhold
    f = tmp_path / "a.xlsx"
    wb.save(f)
    assert read_section.invoke({"path": str(f), "heading": "Sheet"}) == (
        "## Ark: Sheet\n\n| Rad | Navn | Alder |\n| --- | --- | --- |\n| 3 | Kari | 41 |"
    )


def test_small_rows_per_part(tmp_path, monkeypatch):
    monkeypatch.setattr(tables, "ROWS_PER_PART", 2)
    wb = Workbook()
    for row in [["Tall"], [1], [2], [3]]:
        wb.active.append(row)
    f = tmp_path / "a.xlsx"
    wb.save(f)
    outline = document_outline.invoke({"path": str(f)})
    assert "### Rad 2–3" in outline and "### Rad 4–4" in outline


def test_empty_workbook(tmp_path):
    f = tmp_path / "tom.xlsx"
    Workbook().save(f)
    assert "ingen tekst" in read_document.invoke({"path": str(f)})


def test_corrupt_file_returns_error_text(tmp_path):
    bad = tmp_path / "ødelagt.xlsx"
    bad.write_bytes(b"ikke et regneark")
    result = read_document.invoke({"path": str(bad)})
    assert result.startswith("Feil:")
    assert "ikke en gyldig Excel-fil" in result
