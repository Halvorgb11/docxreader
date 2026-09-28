"""Lager samples/tilbud.eml – en e-posttråd med vedlegg.

Kjør:  uv run python scripts/make_sample_eml.py

Inneholder det e-postleseren må håndtere:
- emne og navn med æøå (kodet overskrift, "=?utf-8?...")
- både ren tekst og HTML (leseren skal velge ren tekst)
- sitert tråd i to nivåer: "Den … skrev:" med ">" foran, og inni den et
  Outlook-hode ("-----Opprinnelig melding-----", "Fra:", "Sendt:")
- vedlegg: regneark (.xlsx), Word-dokument (.docx), bilde (.png, kan ikke leses),
  videresendt e-post (message/rfc822) og en innebygd logo som IKKE er et vedlegg
Plantede fakta:
- dypt i tråden (Per): møtet er flyttet til torsdag 8. oktober kl. 10 i rom Fjorden
- i Word-vedlegget: betaling skjer 30 dager etter levering
- i den videresendte e-posten: tilbudet gjelder til 15. oktober 2026
- i regnearket: 4 servere à 45 000 kr fra Nordic Data AS
"""

import io
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from docx import Document
from openpyxl import Workbook

OUT = Path(__file__).parent.parent / "samples" / "tilbud.eml"

PLAIN = """Hei Ola,

Takk for oppsummeringen. Jeg har lagt ved tilbudet fra Nordic Data og et
utkast til avtale. Kan du se over betalingsbetingelsene før fredag?

Hilsen Kari

Den 25. sep. 2026 kl. 14:02 skrev Ola Hansen <ola@fjellbekk.example>:
> Hei Kari,
>
> Vi må bestemme oss for serverleverandør denne uken.
>
> -----Opprinnelig melding-----
> Fra: Per Æsøy <per@fjellbekk.example>
> Sendt: 20. september 2026 09:00
> Til: Ola Hansen
> Emne: Leverandørmøte
>
> Møtet med leverandøren er flyttet til torsdag 8. oktober kl. 10 i rom Fjorden.
"""

HTML = """<html><body><p>Hei Ola,</p><p>Takk for oppsummeringen. Jeg har lagt ved
tilbudet fra Nordic Data og et utkast til avtale.</p><img src="cid:logo"><p>Hilsen Kari</p></body></html>"""


def xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tilbud"
    for row in [("Leverandør", "Vare", "Antall", "Pris per stk"),
                ("Nordic Data AS", "Server R750", 4, 45000),
                ("Nordic Data AS", "Installasjon", 1, 12000)]:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Avtaleutkast", level=0)
    doc.add_heading("Leveranse", level=1)
    doc.add_paragraph("Leverandøren leverer fire servere innen utgangen av november.")
    doc.add_heading("Betalingsbetingelser", level=1)
    doc.add_paragraph("Betaling skjer 30 dager etter levering. Forsinkelsesrente etter loven.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def forwarded_mail() -> EmailMessage:
    m = EmailMessage()
    m["From"] = "Salg <salg@nordicdata.example>"
    m["To"] = "Kari Nordmann <kari@fjellbekk.example>"
    m["Subject"] = "Tilbud 2026-117"
    m["Date"] = "Mon, 21 Sep 2026 08:30:00 +0200"
    m.set_content("Hei,\n\nVedlagt er vårt tilbud. Tilbudet gjelder til 15. oktober 2026.\n\nMvh Nordic Data")
    return m


msg = EmailMessage()
msg["From"] = "Kari Nordmann <kari@fjellbekk.example>"
msg["To"] = "Ola Hansen <ola@fjellbekk.example>"
msg["Cc"] = "Per Æsøy <per@fjellbekk.example>"
msg["Subject"] = "Re: Tilbud på servere – Q4"
msg["Date"] = "Mon, 28 Sep 2026 09:15:00 +0200"
msg["Message-ID"] = "<tilbud-q4@fjellbekk.example>"
msg.set_content(PLAIN)
msg.add_alternative(HTML, subtype="html")
# Innebygd logo: hører til HTML-teksten, ikke et vedlegg.
msg.get_payload()[1].add_related(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64, "image", "png", cid="<logo>")
msg.add_attachment(xlsx_bytes(), maintype="application",
                   subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="tilbud.xlsx")
msg.add_attachment(docx_bytes(), maintype="application",
                   subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="avtaleutkast.docx")
msg.add_attachment(b"\x89PNG\r\n\x1a\n" + b"\x00" * 2048, maintype="image", subtype="png", filename="skisse.png")
msg.add_attachment(forwarded_mail())  # blir message/rfc822

# Faste grenser mellom delene, så filen blir lik hver gang.
for number, part in enumerate(msg.walk()):
    if part.is_multipart():
        part.set_boundary(f"grense-{number}")

OUT.parent.mkdir(exist_ok=True)
OUT.write_bytes(msg.as_bytes(policy=SMTP))
print(f"Lagret {OUT}")
