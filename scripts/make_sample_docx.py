"""Lager samples/prosjektplan.docx – et test-dokument for read_docx.

Kjør:  uv run python scripts/make_sample_docx.py

Dokumentet dekker det read_docx må håndtere:
- tittel og overskrifter på flere nivåer (stilene "Title", "Heading 1", "Heading 2")
- vanlige avsnitt, med fet og kursiv tekst
- punktliste og nummerert liste
- en tabell MIDT i teksten (tester at rekkefølgen bevares)
- topptekst og bunntekst
"""

from pathlib import Path

from docx import Document

OUT = Path(__file__).parent.parent / "samples" / "prosjektplan.docx"

doc = Document()

# Topp- og bunntekst ligger i dokumentets "section", ikke i selve brødteksten.
section = doc.sections[0]
section.header.paragraphs[0].text = "Fjellbekk AS – Internt"
section.footer.paragraphs[0].text = "Konfidensielt"

# Tittel og overskrifter. add_heading(tekst, level) setter stilen for oss:
# level=0 -> "Title", level=1 -> "Heading 1", level=2 -> "Heading 2" osv.
doc.add_heading("Prosjektplan: Ny nettbutikk", level=0)

doc.add_heading("Bakgrunn", level=1)
p = doc.add_paragraph("Omsetningen i Q3 var ")
# Et avsnitt består av "runs": biter av tekst med lik formatering.
p.add_run("4,2 millioner kroner").bold = True
p.add_run(". Vi ønsker å øke netthandelen med ")
p.add_run("minst 20 prosent").italic = True
p.add_run(" i løpet av 2027.")

doc.add_heading("Mål", level=1)
doc.add_paragraph("Lansere ny nettbutikk i november", style="List Bullet")
doc.add_paragraph("Halvere tiden fra bestilling til levering", style="List Bullet")
doc.add_paragraph("Støtte Vipps og faktura", style="List Bullet")

doc.add_heading("Budsjett", level=1)
doc.add_paragraph("Tabellen under viser budsjettet fordelt på poster.")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
header = table.rows[0].cells
header[0].text, header[1].text, header[2].text = "Post", "Beløp (kr)", "Ansvarlig"
for post, belop, ansvarlig in [
    ("Utvikling", "850 000", "Kari Nordmann"),
    ("Design", "200 000", "Ola Hansen"),
    ("Markedsføring", "150 000", "Per Æsøy"),
]:
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = post, belop, ansvarlig

# Dette avsnittet står ETTER tabellen i dokumentet. Hvis read_docx
# skriver det FØR tabellen, har den rotet til rekkefølgen.
doc.add_paragraph("Totalt budsjett er 1,2 millioner kroner.")

doc.add_heading("Fremdrift", level=1)
doc.add_heading("Fase 1: Design", level=2)
doc.add_paragraph("Skisser og brukertester i august.")
doc.add_heading("Fase 2: Utvikling", level=2)
doc.add_paragraph("Sett opp nettbutikkplattform", style="List Number")
doc.add_paragraph("Integrer betaling", style="List Number")
doc.add_paragraph("Test med ekte kunder", style="List Number")

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f"Lagret {OUT}")
