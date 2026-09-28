"""Lager samples/salgsmote.pptx – en liten test-presentasjon.

Kjør:  uv run python scripts/make_sample_pptx.py

Inneholder det PowerPoint-leseren må håndtere:
- tittellysbilde med undertittel
- kulepunkter med innrykk (nivå 0 og 1)
- tabell
- tekstbokser laget i "feil" rekkefølge (nederste først) – tester sortering
- en gruppe av former
- lysbilde uten tittel
- talenotater (der et "skjult" faktum står)
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).parent.parent / "samples" / "salgsmote.pptx"

prs = Presentation()
prs.core_properties.title = "Salgsmøte Q3"
TITLE_SLIDE, TITLE_AND_CONTENT, TITLE_ONLY, BLANK = (prs.slide_layouts[i] for i in (0, 1, 5, 6))

# 1: Tittellysbilde
slide = prs.slides.add_slide(TITLE_SLIDE)
slide.shapes.title.text = "Salgsmøte Q3 2026"
slide.placeholders[1].text = "Region Vest"

# 2: Kulepunkter med innrykk + talenotater
slide = prs.slides.add_slide(TITLE_AND_CONTENT)
slide.shapes.title.text = "Resultater"
body = slide.placeholders[1].text_frame
body.text = "Omsetningen økte med 12 prosent"
for text, level in [("Størst vekst i Bergen", 1), ("Svakere i Stavanger", 1), ("Tre nye storkunder", 0)]:
    p = body.add_paragraph()
    p.text, p.level = text, level
slide.notes_slide.notes_text_frame.text = (
    "Nevn at den største nye kunden er Havbruk Vest AS, med kontrakt på 4,2 MNOK."
)

# 3: Tabell
slide = prs.slides.add_slide(TITLE_ONLY)
slide.shapes.title.text = "Salg per selger"
table = slide.shapes.add_table(4, 3, Inches(1), Inches(1.5), Inches(8), Inches(2)).table
for r, row in enumerate([
    ("Selger", "Salg (MNOK)", "Mål nådd"),
    ("Kari Nordmann", "8,1", "Ja"),
    ("Ola Hansen", "5,4", "Nei"),
    ("Per Æsøy", "6,9", "Ja"),
]):
    for c, value in enumerate(row):
        table.cell(r, c).text = value

# 4: Tekstbokser laget nederst-først, pluss en gruppe
slide = prs.slides.add_slide(TITLE_ONLY)
slide.shapes.title.text = "Neste steg"
slide.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(1)).text_frame.text = "Til slutt: evaluering i desember."
slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = "Først: kickoff med nye kunder i oktober."
group = slide.shapes.add_group_shape()
group.shapes.add_textbox(Inches(1), Inches(3.5), Inches(4), Inches(1)).text_frame.text = "Deretter: opplæring av selgere."

# 5: Lysbilde uten tittel
slide = prs.slides.add_slide(BLANK)
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame.text = "Spørsmål?"

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f"Lagret {OUT}")
