"""Leser for Word-dokumenter (.docx), med python-docx.

En .docx har ingen sider; det bestemmes først når Word viser dokumentet.
Strukturen ligger i stilen til hvert avsnitt ("Heading 1", "List Bullet" …).
"""

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from docxreader.blocks import Block, DocumentError, table_block


def _heading_level(style: str) -> int:
    """Stilnavn -> overskriftsnivå. Word lagrer innebygde stilnavn på engelsk,
    også i norske dokumenter."""
    if style == "Title":
        return 1
    if style.startswith("Heading "):
        number = style.removeprefix("Heading ")
        if number.isdigit():
            return int(number) + 1
    return 0


def _paragraph_to_markdown(paragraph: Paragraph, list_counter: list[int]) -> str:
    """Gjør ett avsnitt om til én linje markdown, basert på stilen.

    `list_counter` er en liste med ett tall som husker hvor langt vi er kommet
    i en nummerert liste. (En liste, fordi et tall ikke kan endres "inne i"
    funksjonen og bli husket utenfor.)
    """
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style is not None else ""

    # Alt som ikke er et nummerert listepunkt, avslutter en nummerert liste.
    if not style.startswith("List Number"):
        list_counter[0] = 0

    if not text:
        return ""
    if level := _heading_level(style):
        return "#" * level + " " + text
    if style.startswith("List Bullet"):
        return f"- {text}"
    if style.startswith("List Number"):
        list_counter[0] += 1
        return f"{list_counter[0]}. {text}"
    return text


def load(path: Path) -> tuple[list[Block], str]:
    """Åpne dokumentet og del det i blokker. Returnerer (blokker, topp-/bunntekst)."""
    try:
        doc = Document(str(path))
    except PackageNotFoundError:
        # python-docx sier "Package not found" også når filen finnes men er
        # ødelagt. Vi gir modellen en tydeligere melding.
        raise DocumentError(f"'{path}' er ikke et gyldig Word-dokument (ødelagt eller feil format).")

    # Topp- og bunntekst ligger i "sections", utenfor brødteksten.
    # set() fjerner duplikater når flere seksjoner har samme tekst.
    headers = {p.text.strip() for s in doc.sections for p in s.header.paragraphs if p.text.strip()}
    footers = {p.text.strip() for s in doc.sections for p in s.footer.paragraphs if p.text.strip()}
    meta = []
    if headers:
        meta.append("Topptekst: " + " / ".join(sorted(headers)))
    if footers:
        meta.append("Bunntekst: " + " / ".join(sorted(footers)))

    # iter_inner_content() gir avsnitt OG tabeller i den rekkefølgen de står i
    # dokumentet. (doc.paragraphs og doc.tables er separate lister og mister
    # rekkefølgen mellom dem.)
    blocks: list[Block] = []
    list_counter = [0]
    prev_was_list = False
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            list_counter[0] = 0
            block = table_block([[cell.text for cell in row.cells] for row in item.rows])
            is_list = False
        else:
            style = item.style.name if item.style is not None else ""
            block = Block(
                _paragraph_to_markdown(item, list_counter),
                len(item.text.split()),
                heading_level=_heading_level(style),
            )
            is_list = style.startswith("List")
        if not block.markdown:
            continue
        # Listepunkter som følger hverandre samles i én blokk (én per linje).
        if is_list and prev_was_list:
            blocks[-1].markdown += "\n" + block.markdown
            blocks[-1].words += block.words
        else:
            blocks.append(block)
        prev_was_list = is_list

    return blocks, "\n\n".join(meta)
