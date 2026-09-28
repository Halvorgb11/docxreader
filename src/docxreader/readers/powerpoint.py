"""Leser for PowerPoint-presentasjoner (.pptx), med python-pptx.

Hvert lysbilde blir én seksjon med overskriften "Lysbilde N: <tittel>", så
innholdsfortegnelse, seksjoner og søk fungerer som for Word.

Et lysbilde består av "shapes" (former): tekstbokser, tabeller, bilder,
diagrammer og grupper av former. De lagres i den rekkefølgen de ble laget,
ikke der de står på lysbildet, så vi sorterer dem etter plassering.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from docxreader.blocks import Block, DocumentError, count_words, heading, table_block

SLIDE_LEVEL = 2  # "## Lysbilde N" – samme nivå som Heading 1 i Word


def _reading_order(shapes) -> list:
    """Sorter former ovenfra og ned, deretter fra venstre mot høyre."""
    return sorted(shapes, key=lambda s: (s.top or 0, s.left or 0))


def _shape_to_blocks(shape) -> list[Block]:
    """Én form -> blokker. Grupper leses rekursivt (en gruppe inneholder former)."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [b for inner in _reading_order(shape.shapes) for b in _shape_to_blocks(inner)]

    if shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        return [table_block(rows)]

    if shape.has_chart:
        chart = shape.chart
        title = chart.chart_title.text_frame.text if chart.has_title else "uten tittel"
        return [Block(f"[Diagram: {title}]", count_words(title))]

    if shape.has_text_frame:
        paragraphs = [(p.level, p.text.strip()) for p in shape.text_frame.paragraphs]
        paragraphs = [(level, text) for level, text in paragraphs if text]
        if not paragraphs:
            return []
        if len(paragraphs) == 1:
            text = paragraphs[0][1]
            return [Block(text, count_words(text))]
        # Flere avsnitt i én tekstboks er som regel kulepunkter.
        # p.level er innrykksnivået (0 = øverst), som vi viser med mellomrom.
        text = "\n".join("  " * level + f"- {t}" for level, t in paragraphs)
        return [Block(text, count_words(text))]

    return []  # bilder, linjer osv. har ingen tekst vi kan lese


def load(path: Path) -> tuple[list[Block], str]:
    try:
        presentation = Presentation(str(path))
    except PackageNotFoundError:
        raise DocumentError(f"'{path}' er ikke en gyldig PowerPoint-fil (ødelagt eller feil format).")

    blocks: list[Block] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title  # None hvis lysbildet ikke har tittelfelt
        title = title_shape.text.strip() if title_shape is not None else ""
        blocks.append(heading(SLIDE_LEVEL, f"Lysbilde {number}: {title}" if title else f"Lysbilde {number}"))

        for shape in _reading_order(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue  # tittelen står allerede i overskriften
            blocks.extend(_shape_to_blocks(shape))

        # Talenotater: det foredragsholderen skal si. Ofte der detaljene står.
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                text = "Talenotater: " + " ".join(notes.split())
                blocks.append(Block(text, count_words(text)))

    meta = f"Presentasjon med {len(presentation.slides)} lysbilder."
    if presentation.core_properties.title:
        meta = f"Tittel: {presentation.core_properties.title}\n" + meta
    return blocks, meta
