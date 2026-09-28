"""Leser for Excel-arbeidsbøker (.xlsx), med openpyxl.

To måter å lese på:
- load_tables(): hvert ark som en Table (rader og kolonner) – brukes av query_table.
- load():        blokker for de vanlige verktøyene. Hvert ark blir én seksjon
                 ("## Ark: Salg") med en markdown-tabell, laget fra load_tables().

Tabellene i et ark finnes med tables.split_tables: titler over tabellen,
flere tabeller under hverandre (adskilt av tomme rader) og fotnoter.
Første rad i hver tabell regnes som overskriftsrad. Hver rad får Excel-radnummeret
sitt i en egen "Rad"-kolonne, så modellen kan si "rad 57" og brukeren finne den.
Store ark deles i biter på tables.ROWS_PER_PART rader ("### Rad 2–101").
Cellekommentarer tas med under tabellen – der står ofte forklaringen på et tall.
"""

import datetime
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from docxreader.blocks import Block, DocumentError, heading
from docxreader.tables import Table, split_tables, table_blocks

SHEET_LEVEL = 2  # "## Ark: navn" – samme nivå som lysbilder og Heading 1


def _format(value, formula) -> str:
    """Én celleverdi -> tekst."""
    if value is None:
        # Formelen har ikke noe lagret resultat (filen er aldri åpnet i Excel
        # etter at formelen ble skrevet). Vis formelen i stedet for ingenting.
        return formula if isinstance(formula, str) and formula.startswith("=") else ""
    if isinstance(value, datetime.datetime) and value.time() == datetime.time(0):
        return value.date().isoformat()  # 2026-09-28, ikke 2026-09-28 00:00:00
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 1200.0 -> 1200
    return str(value)


def _sheet_rows(values_sheet, formulas_sheet) -> list[tuple[int, list[str]]]:
    """Alle ikke-tomme rader som (Excel-radnummer, celletekster).
    Tomme kolonner helt til høyre fjernes."""
    rows = []
    for values, formulas in zip(values_sheet.iter_rows(), formulas_sheet.iter_rows()):
        cells = [_format(v.value, f.value) for v, f in zip(values, formulas)]
        if any(cells):
            rows.append((values[0].row, cells))
    if not rows:
        return []
    # ws.max_column kan være for stor (formatering teller); kutt tomme kolonner.
    width = max(max((i + 1 for i, c in enumerate(cells) if c), default=0) for _, cells in rows)
    return [(number, cells[:width]) for number, cells in rows]


def _sheets(path: Path) -> list[tuple[str, bool, list[Table]]]:
    """Hvert ark som (navn, skjult?, tabeller i arket)."""
    try:
        # To utgaver av samme fil: én med lagrede resultater av formler
        # (data_only=True) og én med selve formlene.
        values_book = load_workbook(path, data_only=True)
        formulas_book = load_workbook(path)
    except (zipfile.BadZipFile, InvalidFileException, KeyError):
        # .xlsx er en zip-fil med XML inni; ødelagt zip eller manglende deler.
        raise DocumentError(f"'{path}' er ikke en gyldig Excel-fil (ødelagt eller feil format).")

    sheets = []
    for values_sheet in values_book.worksheets:
        formulas_sheet = formulas_book[values_sheet.title]
        hidden = values_sheet.sheet_state != "visible"
        tables = split_tables(values_sheet.title, _sheet_rows(values_sheet, formulas_sheet), hidden=hidden)

        # Hver kommentar havner hos tabellen som dekker raden dens (ellers den første).
        for row in formulas_sheet.iter_rows():
            for cell in row:
                if not cell.comment:
                    continue
                text = f"{cell.coordinate}: {' '.join(cell.comment.text.split())}"
                owner = next(
                    (t for t in tables if t.rows and cell.row <= t.rows[-1][0]), tables[0]
                )
                owner.comments.append(text)
        sheets.append((values_sheet.title, hidden, tables))
    return sheets


def load_tables(path: Path) -> list[Table]:
    """Alle tabeller i alle ark – også tomme og skjulte ark."""
    return [table for _, _, tables in _sheets(path) for table in tables]


def _describe(table: Table) -> str:
    return f"{len(table.rows)} rader; kolonner: {', '.join(table.header)}"


def load(path: Path) -> tuple[list[Block], str]:
    blocks: list[Block] = []
    summary = []  # én linje per ark (og tabell), med kolonnenavn – nyttig før query_table
    for name, hidden, tables in _sheets(path):
        if not any(t.header or t.comments for t in tables):
            summary.append(f"- {name}: tomt")
            continue
        title = f"Ark: {name}" + (" (skjult)" if hidden else "")
        if len(tables) == 1:
            summary.append(f"- {name}: {_describe(tables[0])}")
            blocks.extend(table_blocks(tables[0], title, SHEET_LEVEL))
        else:
            # Flere tabeller i samme ark: hver får sin egen underoverskrift.
            summary.append(f"- {name}: {len(tables)} tabeller")
            summary.extend(f"  - {t.name}: {_describe(t)}" for t in tables)
            blocks.append(heading(SHEET_LEVEL, title))
            for number, table in enumerate(tables, start=1):
                label = table.name.removeprefix(f"{name} – ")
                blocks.extend(table_blocks(table, label, SHEET_LEVEL + 1))

    meta = f"Arbeidsbok med {sum(1 for line in summary if line.startswith('- '))} ark:\n" + "\n".join(summary)
    return blocks, meta
