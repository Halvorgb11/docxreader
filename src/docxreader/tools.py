"""Verktøy (tools) som en LangChain-agent kan kalle.

Et verktøy er en vanlig Python-funksjon pyntet med `@tool`. Dekoratoren gjør
funksjonen om til et `BaseTool`-objekt som LangChain kan gi til modellen:

- funksjonsnavnet blir verktøyets navn (`read_docx`),
- docstringen blir beskrivelsen modellen leser for å avgjøre NÅR verktøyet
  skal brukes,
- typehintene (`path: str`) blir et skjema som forteller modellen HVILKE
  argumenter den må sende med.

Modellen kjører aldri koden selv. Den ber om et kall ("kall read_docx med
path='x.docx'"), agenten kjører funksjonen, og returverdien sendes tilbake til
modellen som tekst.

Verktøyene her:
- read_docx          – hele dokumentet (eller innholdsfortegnelsen hvis det er stort)
- docx_outline       – bare innholdsfortegnelsen, med størrelse per seksjon
- read_docx_section  – én seksjon, valgt med overskriften
- search_docx        – finn setninger, listepunkter og tabellrader som inneholder søkeord
"""

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.tools import tool

# Dokumenter (og seksjoner) over dette antallet ord sendes ikke i sin helhet.
# ~3000 ord er omtrent 4000–5000 tokens.
MAX_WORDS = 3000

# search_docx viser høyst så mange treff.
MAX_SEARCH_HITS = 15


# --- Lesing av .docx ------------------------------------------------------
#
# En .docx har ingen sider; det bestemmes først når Word viser dokumentet.
# Strukturen ligger i stilen til hvert avsnitt ("Heading 1", "List Bullet" …).
# Vi deler dokumentet i blokker og gjør hver blokk om til markdown.
#
# Funksjonene og klassene med _ foran er vanlige hjelpere. Bare funksjoner
# med @tool blir verktøy som modellen kan se og kalle.


@dataclass
class _Block:
    """Én del av dokumentet: en overskrift, et avsnitt, en liste eller en tabell."""

    markdown: str
    words: int
    # 0 = ikke en overskrift. 1 = tittel (#), 2 = Heading 1 (##), 3 = Heading 2 (###) …
    heading_level: int = 0
    is_table: bool = False

    @property
    def heading_text(self) -> str:
        return self.markdown.lstrip("#").strip() if self.heading_level else ""


class _DocxError(Exception):
    """En feil som skal vises til modellen som en "Feil: …"-tekst."""


def _heading_level(style: str) -> int:
    """Stilnavn -> overskriftsnivå. Word lagrer innebygde stilnavn på engelsk,
    også i norske dokumenter."""
    if style == "Title":
        return 1
    if style.startswith("Heading "):
        number = style.removeprefix("Heading ")
        if number.isdigit():
            return int(number) + 1
    return 0


def _paragraph_to_markdown(paragraph: Paragraph, list_counter: list[int]) -> str:
    """Gjør ett avsnitt om til én linje markdown, basert på stilen.

    `list_counter` er en liste med ett tall som husker hvor langt vi er kommet
    i en nummerert liste. (En liste, fordi et tall ikke kan endres "inne i"
    funksjonen og bli husket utenfor.)
    """
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style is not None else ""

    # Alt som ikke er et nummerert listepunkt, avslutter en nummerert liste.
    if not style.startswith("List Number"):
        list_counter[0] = 0

    if not text:
        return ""
    if level := _heading_level(style):
        return "#" * level + " " + text
    if style.startswith("List Bullet"):
        return f"- {text}"
    if style.startswith("List Number"):
        list_counter[0] += 1
        return f"{list_counter[0]}. {text}"
    return text


def _table_to_markdown(table: Table) -> str:
    """Gjør en Word-tabell om til en markdown-tabell. Første rad blir overskrift."""

    def clean(cell_text: str) -> str:
        # Linjeskift og | ville ødelagt markdown-tabellen.
        return cell_text.replace("\n", " ").replace("|", "\\|").strip()

    rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""

    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("|" + " --- |" * len(rows[0]))
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _load_docx(path: str) -> tuple[list[_Block], str]:
    """Åpne dokumentet og del det i blokker.

    Returnerer (blokker, topp-/bunntekst som ferdig tekst).
    Kaster _DocxError hvis filen mangler eller ikke kan leses.
    """
    docx_path = Path(path)
    if not docx_path.is_file():
        raise _DocxError(f"fant ingen fil på '{path}'.")
    if docx_path.suffix.lower() != ".docx":
        raise _DocxError(f"'{path}' er ikke en .docx-fil.")
    try:
        doc = Document(str(docx_path))
    except PackageNotFoundError:
        # python-docx sier "Package not found" også når filen finnes men er
        # ødelagt. Vi gir modellen en tydeligere melding.
        raise _DocxError(f"'{path}' er ikke et gyldig Word-dokument (ødelagt eller feil format).")

    # Topp- og bunntekst ligger i "sections", utenfor brødteksten.
    # set() fjerner duplikater når flere seksjoner har samme tekst.
    headers = {p.text.strip() for s in doc.sections for p in s.header.paragraphs if p.text.strip()}
    footers = {p.text.strip() for s in doc.sections for p in s.footer.paragraphs if p.text.strip()}
    meta = []
    if headers:
        meta.append("Topptekst: " + " / ".join(sorted(headers)))
    if footers:
        meta.append("Bunntekst: " + " / ".join(sorted(footers)))

    # iter_inner_content() gir avsnitt OG tabeller i den rekkefølgen de står i
    # dokumentet. (doc.paragraphs og doc.tables er separate lister og mister
    # rekkefølgen mellom dem.)
    blocks: list[_Block] = []
    list_counter = [0]
    prev_was_list = False
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            list_counter[0] = 0
            words = sum(len(cell.text.split()) for row in item.rows for cell in row.cells)
            block = _Block(_table_to_markdown(item), words, is_table=True)
            is_list = False
        else:
            style = item.style.name if item.style is not None else ""
            block = _Block(
                _paragraph_to_markdown(item, list_counter),
                len(item.text.split()),
                heading_level=_heading_level(style),
            )
            is_list = style.startswith("List")
        if not block.markdown:
            continue
        # Listepunkter som følger hverandre samles i én blokk (én per linje).
        if is_list and prev_was_list:
            blocks[-1].markdown += "\n" + block.markdown
            blocks[-1].words += block.words
        else:
            blocks.append(block)
        prev_was_list = is_list

    return blocks, "\n\n".join(meta)


def _render(blocks: list[_Block]) -> str:
    """Blokker -> markdown. Tom linje mellom blokker gir gyldig og lettlest markdown."""
    return "\n\n".join(b.markdown for b in blocks)


def _section_end(blocks: list[_Block], start: int) -> int:
    """En seksjon er overskriften på `start` og alt etter den, fram til neste
    overskrift på samme eller høyere nivå. Returnerer indeksen der den slutter."""
    level = blocks[start].heading_level
    for i in range(start + 1, len(blocks)):
        if 0 < blocks[i].heading_level <= level:
            return i
    return len(blocks)


def _outline(blocks: list[_Block], start: int = 0, end: int | None = None) -> str:
    """Innholdsfortegnelse: hver overskrift med størrelsen på seksjonen sin."""
    end = len(blocks) if end is None else end
    lines = []

    # Tekst før første overskrift ville ellers vært usynlig i oversikten.
    first_heading = next((i for i in range(start, end) if blocks[i].heading_level), end)
    if intro_words := sum(b.words for b in blocks[start:first_heading]):
        lines.append(f"(tekst før første overskrift: ~{intro_words} ord)")

    for i in range(first_heading, end):
        if not blocks[i].heading_level:
            continue
        section = blocks[i : _section_end(blocks, i)]
        words = sum(b.words for b in section)
        tables = sum(b.is_table for b in section)
        size = f"~{words} ord"
        if tables:
            size += f", {tables} tabell" + ("er" if tables > 1 else "")
        lines.append(f"{blocks[i].markdown}  ({size})")
    return "\n".join(lines)


def _heading_paths(blocks: list[_Block]) -> dict[int, list[str]]:
    """For hver overskrift: stien av overskrifter ned til den.

    F.eks. {7: ["Økonomi", "Status"]}. Tittelen (nivå 1) tas ikke med i
    stiene til andre overskrifter, siden den er felles for hele dokumentet.
    """
    stack: list[tuple[int, str]] = []
    paths = {}
    for i, b in enumerate(blocks):
        if not b.heading_level:
            continue
        if b.heading_level == 1:
            paths[i] = [b.heading_text]
            continue
        while stack and stack[-1][0] >= b.heading_level:
            stack.pop()
        stack.append((b.heading_level, b.heading_text))
        paths[i] = [text for _, text in stack]
    return paths


def _find_section(blocks: list[_Block], heading: str) -> int:
    """Finn indeksen til overskriften brukeren (modellen) ba om.

    `heading` kan være en enkelt overskrift ("Budsjett") eller en sti
    ("Økonomi > Status"). Store/små bokstaver spiller ingen rolle.
    Først prøves eksakt treff, deretter treff på deler av teksten.
    """
    wanted = [part.strip().casefold() for part in heading.split(">") if part.strip()]
    if not wanted:
        raise _DocxError("overskriften kan ikke være tom.")
    paths = _heading_paths(blocks)

    def matches(path: list[str], exact: bool) -> bool:
        path = [p.casefold() for p in path]
        last = path[-1] == wanted[-1] if exact else wanted[-1] in path[-1]
        if not last:
            return False
        # De øvrige delene må finnes blant foreldrene, i riktig rekkefølge.
        parents = iter(path[:-1])
        return all(any(w == p for p in parents) for w in wanted[:-1])

    for exact in (True, False):
        hits = [i for i, path in paths.items() if matches(path, exact)]
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            options = "\n".join("- " + " > ".join(paths[i]) for i in hits)
            raise _DocxError(
                f"flere overskrifter passer '{heading}':\n{options}\n"
                "Oppgi hele stien, f.eks. 'Kapittel > Underoverskrift'."
            )

    available = "\n".join("- " + " > ".join(p) for p in paths.values())
    raise _DocxError(f"fant ingen overskrift som passer '{heading}'. Tilgjengelige:\n{available}")


def _block_paths(blocks: list[_Block]) -> list[str]:
    """For hver blokk: overskriftsstien den står under, f.eks. "Økonomi > Status"."""
    heading_paths = _heading_paths(blocks)
    current = "(før første overskrift)"
    result = []
    for i in range(len(blocks)):
        if i in heading_paths:
            current = " > ".join(heading_paths[i])
        result.append(current)
    return result


def _search_units(block: _Block) -> tuple[str, list[str]]:
    """Del en blokk i biter vi søker i, så et treff ikke returnerer et helt avsnitt.

    Returnerer (tekst som vises foran treffene, liste med biter):
    - tabell:  hver rad; tabellens overskriftsrad vises foran, så tallene gir mening
    - liste:   hvert listepunkt
    - avsnitt: hver setning
    """
    if block.is_table:
        lines = block.markdown.splitlines()
        return "\n".join(lines[:2]), lines[2:]
    if "\n" in block.markdown:
        return "", block.markdown.splitlines()
    # Del etter punktum/utrop/spørsmålstegn etterfulgt av mellomrom.
    return "", re.split(r"(?<=[.!?])\s+", block.markdown)


# --- Verktøyene -----------------------------------------------------------

# parse_docstring=True: LangChain leser "Args:"-seksjonen i docstringen og
# legger beskrivelsen av hvert argument inn i skjemaet modellen ser.
@tool(parse_docstring=True)
def read_docx(path: str) -> str:
    """Les et helt Word-dokument (.docx) og returner innholdet som markdown.

    Overskrifter blir #-overskrifter, lister blir markdown-lister og tabeller
    blir markdown-tabeller. Topp- og bunntekst tas med øverst.
    Er dokumentet stort, returneres innholdsfortegnelsen i stedet; bruk da
    read_docx_section for å lese delene du trenger.

    Args:
        path: Filsti til .docx-filen, for eksempel "samples/prosjektplan.docx".
    """
    try:
        blocks, meta = _load_docx(path)
    except _DocxError as e:
        return f"Feil: {e}"
    if not blocks:
        return f"Dokumentet '{path}' inneholder ingen tekst."

    total = sum(b.words for b in blocks)
    if total > MAX_WORDS:
        return (
            f"Dokumentet er stort (~{total} ord), så det returneres ikke i sin helhet.\n"
            "Her er innholdsfortegnelsen. Bruk read_docx_section med en overskrift "
            "herfra for å lese den delen du trenger.\n\n"
            + (meta + "\n\n" if meta else "")
            + _outline(blocks)
        )
    return (meta + "\n\n" if meta else "") + _render(blocks)


@tool(parse_docstring=True)
def docx_outline(path: str) -> str:
    """Vis innholdsfortegnelsen til et Word-dokument (.docx), uten selve teksten.

    Viser alle overskrifter med omtrentlig antall ord og tabeller i hver seksjon.
    Bruk dette først for å finne ut hvor i dokumentet svaret står, og les
    deretter bare den delen med read_docx_section.

    Args:
        path: Filsti til .docx-filen, for eksempel "samples/prosjektplan.docx".
    """
    try:
        blocks, meta = _load_docx(path)
    except _DocxError as e:
        return f"Feil: {e}"

    total = sum(b.words for b in blocks)
    if not any(b.heading_level for b in blocks):
        return f"Dokumentet har ingen overskrifter (~{total} ord totalt). Bruk read_docx for å lese det."
    return (meta + "\n\n" if meta else "") + f"Totalt ~{total} ord.\n\n" + _outline(blocks)


@tool(parse_docstring=True)
def read_docx_section(path: str, heading: str) -> str:
    """Les én seksjon av et Word-dokument (.docx): en overskrift med alt innhold
    under den, inkludert underoverskrifter.

    Finn overskriftene med docx_outline først. Hvis samme overskrift finnes
    flere steder, oppgi stien med " > ", f.eks. "Økonomi > Status".

    Args:
        path: Filsti til .docx-filen, for eksempel "samples/prosjektplan.docx".
        heading: Overskriften til seksjonen, f.eks. "Budsjett" eller "Økonomi > Status". Store/små bokstaver spiller ingen rolle.
    """
    try:
        blocks, _ = _load_docx(path)
        start = _find_section(blocks, heading)
    except _DocxError as e:
        return f"Feil: {e}"

    end = _section_end(blocks, start)
    section = blocks[start:end]
    words = sum(b.words for b in section)
    has_subsections = any(b.heading_level for b in section[1:])

    # Også en seksjon kan være for stor. Har den underoverskrifter, viser vi
    # dem i stedet, så modellen kan gå et nivå dypere.
    if words > MAX_WORDS and has_subsections:
        return (
            f"Seksjonen '{blocks[start].heading_text}' er stor (~{words} ord). "
            "Her er underoverskriftene; les dem enkeltvis med read_docx_section.\n\n"
            + _outline(blocks, start, end)
        )
    return _render(section)


@tool(parse_docstring=True)
def search_docx(path: str, query: str) -> str:
    """Søk etter ord i et Word-dokument (.docx) og få bare de treffende
    setningene, listepunktene og tabellradene, med overskriften de står under.

    Alle ordene i søket må finnes i samme setning/rad. Store/små bokstaver
    spiller ingen rolle, og deler av ord gir treff ("sikkerhet" finner
    "sikkerhetshendelser"). Bruk dette når du leter etter noe bestemt og ikke
    vet hvilken seksjon det står i. Les mer sammenheng med read_docx_section.

    Args:
        path: Filsti til .docx-filen, for eksempel "samples/prosjektplan.docx".
        query: Ett eller flere søkeord, f.eks. "driftsresultat" eller "kontaktperson nord". Korte ord og ordstammer gir flest treff.
    """
    words = query.casefold().split()
    if not words:
        return "Feil: søket kan ikke være tomt."
    try:
        blocks, _ = _load_docx(path)
    except _DocxError as e:
        return f"Feil: {e}"

    # Samle treff: (blokk, overskriftssti, tekst foran treffene, treffende biter).
    hits = []
    for block, block_path in zip(blocks, _block_paths(blocks)):
        prefix, units = _search_units(block)
        matched = [u for u in units if all(w in u.casefold() for w in words)]
        if matched:
            hits.append((block, block_path, prefix, matched))

    total = sum(len(matched) for *_, matched in hits)
    if total == 0:
        return (
            f"Ingen treff for '{query}'. Prøv færre eller kortere ord (en ordstamme "
            "finner også sammensatte ord), eller se innholdsfortegnelsen med docx_outline."
        )

    lines = [f"{total} treff for '{query}':"]
    shown = 0
    previous_path = None
    for block, block_path, prefix, matched in hits:
        if shown >= MAX_SEARCH_HITS:
            break
        matched = matched[: MAX_SEARCH_HITS - shown]
        shown += len(matched)
        # Treff under samme overskrift grupperes under én [sti]-linje.
        if block_path != previous_path:
            lines.append(f"\n[{block_path}]")
            previous_path = block_path
        if prefix:  # tabell: overskriftsraden først, så tallene gir mening
            lines.append(prefix)
        # Setninger fra vanlige avsnitt får "- " foran. Tabellrader, listepunkter
        # og overskrifter har allerede sin egen markdown-form.
        is_plain_paragraph = not (block.is_table or block.heading_level or "\n" in block.markdown)
        lines.extend(f"- {u}" if is_plain_paragraph else u for u in matched)

    if shown < total:
        lines.append(
            f"\nViser {shown} av {total} treff. Gjør søket mer presist (flere ord), "
            "eller les en seksjon med read_docx_section."
        )
    return "\n".join(lines)
